from pathlib import Path

import openpyxl


class ExcelReader:
    def __init__(self, file_path, sheet_name=None):
        path = Path(file_path)
        if not path.is_absolute():
            base_dir = Path(__file__).resolve().parent.parent.parent.parent / "testdata"
            path = base_dir / file_path

        if not path.exists():
            raise FileNotFoundError(f"文件不存在: {path}")

        self.file_path = path
        self.workbook = openpyxl.load_workbook(self.file_path)

        if sheet_name is None:
            self.sheet = self.workbook.worksheets[0]
        else:
            if sheet_name not in self.workbook.sheetnames:
                raise ValueError(
                    f"Sheet '{sheet_name}' 不存在，可用的Sheet: {self.workbook.sheetnames}"
                )
            self.sheet = self.workbook[sheet_name]

    def __enter__(self):
        return self

    def __exit__(self, exc_type, exc_val, exc_tb):
        self.close()

    def get_all_sheet_names(self):
        return self.workbook.sheetnames

    def get_sheet_name(self):
        return self.sheet.title

    def get_headers(self):
        headers = []
        for cell in self.sheet[1]:
            if cell.value is not None:
                headers.append(cell.value)
        return headers

    def get_data(self):
        headers = self.get_headers()
        data_list = []

        for row in self.sheet.iter_rows(min_row=2, max_row=self.sheet.max_row,
                                        max_col=len(headers)):
            row_values = [cell.value for cell in row]
            if all(v is None for v in row_values):
                continue
            row_dict = dict(zip(headers, row_values))
            data_list.append(row_dict)

        return data_list

    def get_column_data(self, column_name):
        headers = self.get_headers()
        if column_name not in headers:
            raise ValueError(f"字段名 '{column_name}' 不存在，可用字段: {headers}")

        col_index = headers.index(column_name) + 1
        column_data = []

        for row in self.sheet.iter_rows(min_row=2, max_row=self.sheet.max_row,
                                        min_col=col_index, max_col=col_index):
            value = row[0].value
            if value is not None:
                column_data.append(value)

        return column_data

    def close(self):
        self.workbook.close()
